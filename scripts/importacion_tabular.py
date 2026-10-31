"""Local tabular readers: literal text, explicit date mapping, no clinical inference."""

import hashlib
import math
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any


class TableImportError(ValueError):
    """Safe structural diagnostics, without source identifiers or clinical text."""


@dataclass(frozen=True)
class SourceRecord:
    source_id: str
    text: str
    document_date: date | None
    row_number: int


def parse_document_date(value: object, text_format: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            parsed = datetime.strptime(value.strip(), text_format).date()
            # Python's %y convention must not silently decide the century.
            if "%y" in text_format:
                raise TableImportError(
                    "Para fechas textuales usa año de cuatro cifras o una política explícita."
                )
            return parsed
        except ValueError:
            raise TableImportError(
                "Fecha inválida para el formato declarado; revisa la fila de origen."
            ) from None
    raise TableImportError(
        "La fecha debe ser una fecha Excel reconocida o texto con formato declarado."
    )


def parse_xlsx(
    path: Path,
    id_column: str,
    text_column: str,
    *,
    sheet: str | None,
    date_column: str | None,
    date_format: str = "%d/%m/%Y",
) -> tuple[list[SourceRecord], str, dict[str, Any]]:
    try:
        import openpyxl
    except ImportError:
        raise TableImportError(
            "El formato XLSX requiere el extra lab: pip install '.[lab]'."
        ) from None
    import warnings

    workbook = None
    try:
        # External links and formulas are not evaluated. Suppress library warnings that may
        # disclose arbitrary source metadata; report only our safe structural messages.
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = openpyxl.load_workbook(
                path, read_only=True, data_only=False, keep_links=False
            )
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                raise TableImportError("Declara hoja_entrada cuando el libro tenga varias hojas.")
            worksheet = workbook[workbook.sheetnames[0]]
        elif sheet not in workbook.sheetnames:
            raise TableImportError("La hoja configurada no existe en el Excel.")
        else:
            worksheet = workbook[sheet]
        rows = worksheet.iter_rows()
        first = next(rows, ())
        header = [cell.value for cell in first]
        required = [id_column, text_column] + ([date_column] if date_column else [])
        if any(header.count(name) != 1 for name in required):
            raise TableImportError("Faltan columnas requeridas o están duplicadas en la hoja.")
        id_index, text_index = header.index(id_column), header.index(text_column)
        date_index = header.index(date_column) if date_column else None
        records = []
        empty_rows = 0
        for row_number, row in enumerate(rows, 2):
            if all(cell.value is None for cell in row):
                empty_rows += 1
                continue
            selected = [row[header.index(name)] for name in required]
            if any(cell.data_type in ("f", "e") for cell in selected):
                raise TableImportError(
                    f"Fórmula o error Excel en una columna requerida, fila {row_number}."
                )
            source_id, text = row[id_index].value, row[text_index].value
            if isinstance(source_id, bool) or not isinstance(source_id, str | int | float):
                raise TableImportError(f"Identificador inválido en fila {row_number}.")
            if isinstance(source_id, float):
                if not math.isfinite(source_id) or not source_id.is_integer():
                    raise TableImportError(
                        f"Identificador numérico no entero en fila {row_number}."
                    )
                source_id = int(source_id)
            source_id = str(source_id)
            if not source_id.strip() or not isinstance(text, str) or not text.strip():
                raise TableImportError(
                    f"Identificador o texto vacío/no textual en fila {row_number}."
                )
            try:
                document_date = (
                    parse_document_date(row[date_index].value, date_format)
                    if date_index is not None
                    else None
                )
            except TableImportError:
                raise TableImportError(
                    f"Fecha inválida o ambigua en fila {row_number}; revisa el mapeo declarado."
                ) from None
            records.append(SourceRecord(source_id, text, document_date, row_number))
        if not records:
            raise TableImportError("La hoja no contiene registros documentales.")
        return (
            records,
            hashlib.sha256(path.read_bytes()).hexdigest(),
            {
                "hoja": worksheet.title,
                "filas_vacias_omitidas": empty_rows,
                "fechas_presentes": sum(record.document_date is not None for record in records),
                "formulas_evaluadas": False,
            },
        )
    except TableImportError:
        raise
    except Exception:
        raise TableImportError(
            "No se pudo leer el XLSX; revisa formato, columnas y valores."
        ) from None
    finally:
        if workbook is not None:
            workbook.close()
